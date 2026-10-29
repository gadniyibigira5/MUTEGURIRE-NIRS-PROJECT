#!/usr/bin/env python3
"""Heuristically extract EICV7 district poverty figures from PowerPoint files."""
from __future__ import annotations
import argparse, json, re
from pathlib import Path
from pptx import Presentation

DISTRICTS = ["Nyarugenge", "Gasabo", "Kicukiro", "Nyanza", "Gisagara", "Nyaruguru", "Huye", "Nyamagabe", "Ruhango", "Muhanga", "Kamonyi", "Karongi", "Rutsiro", "Rubavu", "Nyabihu", "Ngororero", "Rusizi", "Nyamasheke", "Rulindo", "Gakenke", "Musanze", "Burera", "Gicumbi", "Rwamagana", "Nyagatare", "Gatsibo", "Kayonza", "Kirehe", "Ngoma", "Bugesera"]

def parse_number(text: str) -> float | None:
    """Parse a single percentage-like numeric cell (with or without a percent sign)."""
    cleaned = str(text or "").strip().replace("\u00a0", " ").replace(",", ".")
    match = re.fullmatch(r"\(?\s*(\d{1,2}(?:\.\d+)?)\s*%?\s*\)?", cleaned)
    if not match: return None
    value = float(match.group(1))
    return value if 0 <= value <= 100 else None

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx-dir", required=True, help="Folder containing manually downloaded EICV7 district PPTX files")
    parser.add_argument("--out", default="tools/out/eicv_extracted.json", help="Output JSON path")
    args = parser.parse_args()
    root = Path(args.pptx_dir)
    if not root.is_dir(): parser.error(f"PowerPoint folder not found: {root}")
    files = sorted(p for p in root.rglob("*.pptx") if not p.name.startswith("~$"))
    result: dict[str, dict] = {}; scanned = 0
    for path in files:
        try: presentation = Presentation(str(path))
        except Exception as exc:
            print(f"WARN: could not read {path}: {exc}"); continue
        scanned += 1
        for slide_no, slide in enumerate(presentation.slides, start=1):
            shape_texts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False): shape_texts.append(shape.text_frame.text)
                if not getattr(shape, "has_table", False): continue
                table = shape.table
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                header_text = " ".join(" ".join(row) for row in rows[:3])
                poverty_table = bool(re.search(r"poverty|poor|ubukene", header_text, re.I))
                if not poverty_table: continue
                for row in rows:
                    if not row: continue
                    district = next((name for name in DISTRICTS if row[0].strip().casefold() == name.casefold()), None)
                    if not district: continue
                    candidates = [parse_number(cell) for cell in row[1:]]
                    candidates = [v for v in candidates if v is not None]
                    if len(candidates) != 1:
                        print(f"WARN: {district} in {path.name} slide {slide_no} has {len(candidates)} numeric poverty candidates; skipped for manual review")
                        continue
                    if district not in result:
                        result[district] = {"pov": {"value": candidates[0], "slide": slide_no, "file": path.name}}
            # Also support slides that present one district as text rather than a table.
            slide_text = "\n".join(shape_texts)
            if not re.search(r"poverty|poor|ubukene", slide_text, re.I): continue
            for name in DISTRICTS:
                if name in result or not re.search(rf"(?<![A-Za-z]){re.escape(name)}(?![A-Za-z])", slide_text, re.I): continue
                values = [float(x.replace(",", ".")) for x in re.findall(r"(?<!\w)(\d{1,2}(?:[.,]\d+)?)\s*%", slide_text)]
                if len(values) == 1 and 0 <= values[0] <= 100:
                    result[name] = {"pov": {"value": values[0], "slide": slide_no, "file": path.name}}
                elif values:
                    print(f"WARN: {name} in {path.name} slide {slide_no} has {len(values)} candidate percentages; skipped for manual review")
    out = Path(args.out); out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    missing = [name for name in DISTRICTS if name not in result]
    print(f"PowerPoint files scanned: {scanned}; districts matched: {len(result)}/30; output: {out}")
    if missing:
        print("Districts with no unambiguous match:")
        for name in missing: print(f"  - {name}")
    print("REVIEW REQUIRED: check the exact indicator definition, year, source slide, and whether each percentage is a district poverty rate.")
    return 0

if __name__ == "__main__": raise SystemExit(main())
